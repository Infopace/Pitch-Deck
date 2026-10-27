"""Guards the template_data -> PPTX paths that silently degraded to bullets.

`proof.cases`, `galleryS.slots` and `persona.jobs/pains` all had no path
through the python-pptx renderer: the builders sniffed `td` for shapes they
didn't recognise and fell through to `_title_bullets`, so the slide rendered
empty. Nothing failed — the deck just quietly lost a slide's worth of content,
which is exactly the kind of regression a unit test on the registry cannot see.

These tests render a real .pptx in memory and assert the content is on it.
"""
import asyncio
import io

import pytest
from pptx import Presentation

from app.agents.content.content_agent import DeckContent, SlideContent
from app.decks.registry import resolve_sections
from app.decks.sections import SECTION_SPECS
from app.ppt.engine.renderer import PPTRenderer
from app.services.branding.extractor import BrandingData


def _render(deck_type: str, template_data: dict) -> str:
    """Render the deck type's full slide sequence and return all of its text."""
    sections = resolve_sections(deck_type)
    slides = [
        SlideContent(
            slide_number=i,
            slide_type=key,
            layout=SECTION_SPECS[key].layout,
            title=SECTION_SPECS[key].label,
        )
        for i, key in enumerate(sections, start=1)
    ]
    deck = DeckContent(slides=slides, deck_title="Acme", deck_subtitle="")
    data = asyncio.run(
        PPTRenderer().render(
            deck,
            BrandingData(company_name="Acme"),
            template_data=template_data,
            sections=sections,
        )
    )
    prs = Presentation(io.BytesIO(data))
    return "\n".join(
        sh.text_frame.text
        for slide in prs.slides
        for sh in slide.shapes
        if sh.has_text_frame
    )


PROOF_TD = {
    "eyebrow": "Sales Presentation",
    "proof": {
        "headline": "Proof",
        "cases": [
            {
                "c": "Northwind Logistics",
                "m": "-42%",
                "ml": "support cost per ticket",
                "d": "Was 9.1 min average handle time, now 5.3.",
                "q": "It paid for itself in a quarter.",
            }
        ],
        "logos": ["Northwind"],
    },
}

PERSONA_GALLERY_TD = {
    "eyebrow": "Product Overview",
    "persona": {
        "headline": "Who it's for",
        "who": "Support leads running 10-40 agents",
        "context": "Monday triage, inbox at 400+",
        "jobs": [{"k": "Triage fast", "t": "clear the overnight queue before standup"}],
        "pains": [{"k": "Manual tagging", "t": "every ticket hand-routed"}],
        "quote": "I lose an hour a day to routing.",
    },
    "galleryS": {
        "headline": "Product",
        "slots": [
            {"id": "g1", "ph": "Dashboard screenshot", "span": True},
            {"id": "g2", "ph": "Mobile app"},
            {"id": "g3", "ph": "Product in use"},
            {"id": "g4", "ph": "Customer / press"},
            {"id": "g5", "ph": "Reporting view", "span": True},
        ],
    },
}


@pytest.mark.parametrize(
    "needle",
    ["Northwind Logistics", "-42%", "support cost per ticket", "paid for itself"],
)
def test_proof_cases_render(needle):
    assert needle in _render("sales", PROOF_TD)


@pytest.mark.parametrize(
    "needle",
    ["Support leads running 10-40 agents", "Jobs to Be Done", "Triage fast", "Manual tagging"],
)
def test_persona_renders(needle):
    assert needle in _render("product", PERSONA_GALLERY_TD)


@pytest.mark.parametrize("needle", ["Dashboard screenshot", "Reporting view", "Mobile app"])
def test_gallery_slot_captions_render(needle):
    assert needle in _render("product", PERSONA_GALLERY_TD)


def test_footer_uses_deck_type_eyebrow_not_investor():
    """Every deck type used to be footed 'Investor Presentation'."""
    text = _render("sales", PROOF_TD)
    assert "Sales Presentation" in text
    assert "Investor Presentation" not in text


def test_investor_footer_still_correct():
    """Regression: the deck type that was already right must stay right."""
    text = _render("investor", {"eyebrow": "Investor Presentation"})
    assert "Investor Presentation" in text


# ── Gallery photos ────────────────────────────────────────────────────────────
# Uploaded photos live in projects.assets but had no path into render() at all,
# so a downloaded deck showed captioned placeholders where the on-screen deck
# showed the real images.

def _png(w: int, h: int) -> bytes:
    """A solid-colour PNG of a given size, for aspect-ratio/crop assertions."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (90, 120, 220)).save(buf, format="PNG")
    return buf.getvalue()


def _render_pptx(deck_type: str, template_data: dict, **kw):
    sections = resolve_sections(deck_type)
    slides = [
        SlideContent(slide_number=i, slide_type=key,
                     layout=SECTION_SPECS[key].layout, title=SECTION_SPECS[key].label)
        for i, key in enumerate(sections, start=1)
    ]
    data = asyncio.run(PPTRenderer().render(
        DeckContent(slides=slides, deck_title="Acme", deck_subtitle=""),
        BrandingData(company_name="Acme"),
        template_data=template_data, sections=sections, **kw))
    return Presentation(io.BytesIO(data))


def _pictures(prs):
    return [sh for slide in prs.slides for sh in slide.shapes if sh.shape_type == 13]


def test_gallery_photos_are_placed_as_pictures():
    prs = _render_pptx("product", PERSONA_GALLERY_TD,
                       gallery_bytes={"g1": _png(800, 600), "g3": _png(400, 400)})
    assert len(_pictures(prs)) == 2


def test_gallery_photo_is_cropped_not_stretched():
    """A wide photo in a near-square tile must be side-cropped, never squashed."""
    prs = _render_pptx("product", PERSONA_GALLERY_TD, gallery_bytes={"g2": _png(2000, 500)})
    pic = _pictures(prs)[0]
    assert pic.crop_left > 0 and pic.crop_left == pic.crop_right, "should crop evenly on x"
    assert pic.crop_top == 0 and pic.crop_bottom == 0


def test_slots_without_a_photo_keep_their_caption():
    prs = _render_pptx("product", PERSONA_GALLERY_TD, gallery_bytes={"g1": _png(800, 600)})
    text = "\n".join(sh.text_frame.text for s in prs.slides
                     for sh in s.shapes if sh.has_text_frame)
    assert "Mobile app" in text and "Reporting view" in text


def test_undecodable_bytes_fall_back_to_the_caption():
    """An SVG (or anything Pillow can't open) must not break the slide."""
    prs = _render_pptx("product", PERSONA_GALLERY_TD, gallery_bytes={"g1": b"<svg/>"})
    text = "\n".join(sh.text_frame.text for s in prs.slides
                     for sh in s.shapes if sh.has_text_frame)
    assert _pictures(prs) == []
    assert "Dashboard screenshot" in text


def test_gallery_still_renders_with_no_photos():
    prs = _render_pptx("product", PERSONA_GALLERY_TD)
    text = "\n".join(sh.text_frame.text for s in prs.slides
                     for sh in s.shapes if sh.has_text_frame)
    assert "Dashboard screenshot" in text


def test_market_chart_survives_a_non_dict_data_point():
    """Production regression (2026-10-27): a sparse idea-stage deck had no
    template_data.market, so _market_sizing fell back to SlideContent.data_points
    — whose entries the model's JSON output does not guarantee are dicts. One
    plain string among them crashed the whole render with
    AttributeError: 'str' object has no attribute 'get', failing deck
    generation entirely rather than just degrading that one chart."""
    deck_type = "investor"
    sections = resolve_sections(deck_type)
    slides = [
        SlideContent(slide_number=i, slide_type=key,
                     layout=SECTION_SPECS[key].layout, title=SECTION_SPECS[key].label,
                     data_points=(
                         # No template_data.market -> falls back to exactly this list.
                         [{"l": "TAM", "v": "$47B"}, "23% CAGR", {"l": "SOM", "v": "$1.2B"}]
                         if key == "market" else []
                     ))
        for i, key in enumerate(sections, start=1)
    ]
    deck = DeckContent(slides=slides, deck_title="Acme", deck_subtitle="")
    data = asyncio.run(
        PPTRenderer().render(deck, BrandingData(company_name="Acme"),
                              template_data={}, sections=sections)
    )
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(sections)
